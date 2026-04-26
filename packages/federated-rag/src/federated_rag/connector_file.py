"""File connector for local/mounted file system data sources.

Implements BaseConnector for directories containing PDF, DOCX, CSV, TXT, and
MD files. Extracts text content and converts to LightRAG custom KG format.

See: Architecture Plan Section 4.1 -- Base Connector (FileConnector).
"""

from __future__ import annotations

import csv
import hashlib
import io
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, AsyncIterator, Optional

from federated_rag.connector_base import BaseConnector
from federated_rag.models import (
    ColumnSchema,
    ConnectorType,
    InformationSchema,
    TableSchema,
)

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".csv", ".txt", ".md"})


def _sha1_short(text: str) -> str:
    return hashlib.sha1(text.encode()).hexdigest()[:12]


def _extract_text_from_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n\n".join(pages).strip()
    if not text:
        raise ValueError(f"No text extracted from PDF: {path}")
    return text


def _extract_text_from_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    table_cells = []
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_cells.append(" | ".join(cells))
    return "\n".join(paragraphs + table_cells).strip()


def _extract_text_from_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    """Return (headers, rows_as_dicts)."""
    text = path.read_text(encoding="utf-8", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    headers = list(reader.fieldnames or [])
    rows = list(reader)
    return headers, rows


def _extract_text_from_plain(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace").strip()


class FileConnector(BaseConnector):
    """Connector for local file system directories."""

    connector_type = ConnectorType.FILE

    def __init__(self, source_id: Optional[str] = None) -> None:
        self._source_id = source_id or f"file_{uuid.uuid4().hex[:8]}"
        self._root: Optional[Path] = None
        self._files: list[Path] = []

    async def connect(self, config: dict[str, Any]) -> bool:
        path = Path(config["path"])
        if not path.exists():
            return False
        self._root = path
        self._files = self._scan_files()
        return True

    async def disconnect(self) -> None:
        self._root = None
        self._files = []

    async def test_connection(self) -> bool:
        return self._root is not None and self._root.exists()

    async def extract_schema(self) -> InformationSchema:
        if self._root is None:
            raise RuntimeError("Not connected. Call connect() first.")

        tables: list[TableSchema] = []
        for f in self._files:
            if f.suffix.lower() == ".csv":
                try:
                    headers, rows = _extract_text_from_csv(f)
                    columns = [
                        ColumnSchema(
                            name=h,
                            data_type="text",
                            nullable=True,
                            is_primary_key=False,
                            is_foreign_key=False,
                        )
                        for h in headers
                    ]
                    tables.append(
                        TableSchema(name=f.name, columns=columns, row_count=len(rows))
                    )
                except Exception:
                    tables.append(TableSchema(name=f.name, columns=[]))
            else:
                tables.append(
                    TableSchema(
                        name=f.name,
                        columns=[
                            ColumnSchema(
                                name="content",
                                data_type="text",
                                nullable=False,
                                is_primary_key=False,
                                is_foreign_key=False,
                            )
                        ],
                    )
                )

        return InformationSchema(
            source_id=self._source_id,
            source_type=ConnectorType.FILE,
            tables=tables,
            raw_schema={"file_count": len(self._files)},
            extracted_at=datetime.now(tz=timezone.utc),
        )

    async def extract_content(
        self,
        resource_filter: Optional[list[str]] = None,
        since: Optional[datetime] = None,
    ) -> AsyncIterator[dict[str, Any]]:
        if self._root is None:
            raise RuntimeError("Not connected. Call connect() first.")

        files = self._files
        if resource_filter:
            files = [f for f in files if f.name in resource_filter]
        if since:
            files = [
                f for f in files
                if datetime.fromtimestamp(f.stat().st_mtime, tz=timezone.utc) > since
            ]

        for f in files:
            try:
                kg = self._file_to_kg(f)
                if kg["chunks"]:
                    yield kg
            except Exception:
                continue

    async def list_resources(self) -> list[dict[str, Any]]:
        if self._root is None:
            raise RuntimeError("Not connected. Call connect() first.")
        return [
            {
                "name": f.name,
                "type": f.suffix.lstrip("."),
                "size_bytes": f.stat().st_size,
                "path": str(f),
            }
            for f in self._files
        ]

    # -- internals --

    def _scan_files(self) -> list[Path]:
        if self._root is None:
            return []
        if self._root.is_file():
            if self._root.suffix.lower() in SUPPORTED_EXTENSIONS:
                return [self._root]
            return []
        return sorted(
            f
            for f in self._root.rglob("*")
            if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS
        )

    def _file_to_kg(self, f: Path) -> dict[str, Any]:
        suffix = f.suffix.lower()
        doc_id = f"doc:{f.name}:{_sha1_short(str(f))}"

        if suffix == ".csv":
            return self._csv_to_kg(f, doc_id)

        if suffix == ".pdf":
            text = _extract_text_from_pdf(f)
        elif suffix == ".docx":
            text = _extract_text_from_docx(f)
        else:
            text = _extract_text_from_plain(f)

        if not text:
            return {"chunks": [], "entities": [], "relationships": []}

        source_id = f"{self._source_id}:{doc_id}"
        return {
            "chunks": [
                {
                    "source_id": source_id,
                    "content": text,
                    "file_path": f"{self._source_id}/{f.name}",
                }
            ],
            "entities": [
                {
                    "entity_name": f"Document:{f.name}",
                    "entity_type": "Document",
                    "description": f"File: {f.name} ({suffix.lstrip('.')} format, {f.stat().st_size} bytes)",
                    "source_id": source_id,
                    "file_path": f"{self._source_id}/{f.name}",
                }
            ],
            "relationships": [],
        }

    def _csv_to_kg(self, f: Path, doc_id: str) -> dict[str, Any]:
        headers, rows = _extract_text_from_csv(f)
        chunks = []
        entities = []

        for i, row in enumerate(rows):
            source_id = f"{self._source_id}:{doc_id}:row_{i}"
            content_parts = [f"{k}={v}" for k, v in row.items() if v]
            content = f"Row from {f.name}: {', '.join(content_parts)}"
            chunks.append({
                "source_id": source_id,
                "content": content,
                "file_path": f"{self._source_id}/{f.name}",
            })
            entities.append({
                "entity_name": f"CSVRow:{f.name}:{i}",
                "entity_type": "CSVRow",
                "description": content,
                "source_id": source_id,
                "file_path": f"{self._source_id}/{f.name}",
            })

        return {"chunks": chunks, "entities": entities, "relationships": []}
